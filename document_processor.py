import os
import re

import pymupdf
from docx import Document
from pptx import Presentation


# ==========================================
# 1. CLEAN TEXT
# ==========================================

def clean_line(line):
    if not line:
        return ""

    # Remove invisible characters
    line = line.replace("\u200b", "")
    line = line.replace("\ufeff", "")
    line = line.replace("\xa0", " ")

    # Normalize whitespace
    line = re.sub(r"\s+", " ", line)

    return line.strip()


# ==========================================
# 2. EXTRACT TEXT FROM PDF
# ==========================================

def extract_text_from_pdf(file_path):

    doc = pymupdf.open(file_path)

    full_text = ""

    for page in doc:
        full_text += page.get_text() + "\n"

    doc.close()

    return full_text


# ==========================================
# 3. EXTRACT TEXT FROM DOCX
# ==========================================

def extract_text_from_docx(file_path):

    document = Document(file_path)

    lines = []

    # Read paragraphs
    for paragraph in document.paragraphs:

        text = clean_line(
            paragraph.text
        )

        if text:
            lines.append(text)

    # Read tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:

                text = clean_line(
                    cell.text
                )

                if text:
                    row_text.append(text)

            if row_text:

                lines.append(
                    " | ".join(row_text)
                )

    return "\n".join(lines)


# ==========================================
# 4. EXTRACT TEXT FROM PPTX
# ==========================================

def extract_text_from_pptx(file_path):

    presentation = Presentation(
        file_path
    )

    lines = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        lines.append(
            f"Slide {slide_number}"
        )

        for shape in slide.shapes:

            if not hasattr(
                shape,
                "text"
            ):
                continue

            text = clean_line(
                shape.text
            )

            if text:
                lines.append(text)

    return "\n".join(lines)


# ==========================================
# 5. EXTRACT TEXT FROM ANY SUPPORTED FILE
# ==========================================

def extract_text_from_file(file_path):

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    if extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    if extension == ".pptx":

        return extract_text_from_pptx(
            file_path
        )

    raise ValueError(
        "Unsupported file format. "
        "Use PDF, DOCX, or PPTX."
    )


# ==========================================
# 6. NORMALIZE CHAPTER NAME
# ==========================================

def normalize_chapter_name(chapter):

    chapter = clean_line(chapter)

    if not chapter:
        return "General"

    match = re.match(
        r"^Chapter\s+(\d+)\s*[:.]?\s*(.+)$",
        chapter,
        re.IGNORECASE
    )

    if not match:
        return chapter

    number = int(
        match.group(1)
    )

    title = match.group(2).strip()

    title = re.sub(
        r"\s+",
        " ",
        title
    )

    # Remove "in Python" at the end
    title = re.sub(
        r"\s+in\s+Python$",
        "",
        title,
        flags=re.IGNORECASE
    )

    # Fix common extraction issues
    title = title.replace(
        "ControlFlow",
        "Control Flow"
    )

    title = re.sub(
        r"Data\s+Structuress+",
        "Data Structures",
        title,
        flags=re.IGNORECASE
    )

    title = title.rstrip(
        " .:-"
    )

    # Canonical Python chapter names
    canonical_titles = {

        1:
            "Introduction to Python",

        2:
            "Environment Setup",

        3:
            "Python Syntax Basics",

        4:
            "Data Types",

        5:
            "Operators",

        6:
            "String Operations",

        7:
            "Control Flow",

        8:
            "Loops and Iteration",

        9:
            "Data Structures",

        10:
            "Functions",

        11:
            "Advanced Functions",

        12:
            "Scope and Namespaces",

        13:
            "Modules and Packages",

        14:
            "Object-Oriented Programming",

        15:
            "Advanced OOP",

        16:
            "File Handling and Error Management",

        17:
            "Advanced Python Concepts",

        18:
            "Concurrent and Asynchronous Programming"
    }

    if number in canonical_titles:

        title = canonical_titles[number]

    return (
        f"Chapter {number}: {title}"
    )


# ==========================================
# 7. DOCUMENT TYPE
# ==========================================

def get_document_type(file_path):

    filename = os.path.basename(
        file_path
    ).lower()

    question_bank_keywords = [

        "question",
        "question bank",
        "questionbank",
        "previous year",
        "previousyear",
        "pyq",
        "paper",
        "exam",
        "questionpaper",
        "question paper"
    ]

    # Question bank must be checked first
    for keyword in question_bank_keywords:

        if keyword in filename:

            return "question_bank"

    # C programming files
    if (
        "c_programming" in filename
        or "c programming" in filename
        or "c basics" in filename
    ):

        return "simple_number_style"

    # Python files
    if "python" in filename:

        return "chapter_style"

    return "general"


# ==========================================
# 8. DETECT CHAPTER
# ==========================================

def detect_chapter(
    line,
    document_type
):

    line = clean_line(line)

    if not line:
        return None

    # --------------------------------------
    # Python chapter style
    # --------------------------------------

    if document_type == "chapter_style":

        match = re.match(
            r"^Chapter\s+\d+\s*[:.]?\s*.+$",
            line,
            re.IGNORECASE
        )

        if match:

            return normalize_chapter_name(
                line
            )

        return None

    # --------------------------------------
    # Simple numbered chapter style
    # --------------------------------------

    if document_type == "simple_number_style":

        match = re.match(
            r"^\d+\.\s+[A-Za-z][A-Za-z0-9 &()/+\\\-'.:]*$",
            line
        )

        if match:

            return line

        return None

    return None


# ==========================================
# 9. SPLIT INTO CHAPTERS
# ==========================================

def split_into_chapters(
    text,
    default_chapter="General",
    document_type="chapter_style"
):

    lines = []

    for line in text.splitlines():

        cleaned = clean_line(line)

        if cleaned:
            lines.append(cleaned)

    sections = []

    current_chapter = default_chapter

    current_lines = []

    found_first_heading = False

    for line in lines:

        heading = detect_chapter(
            line,
            document_type
        )

        if heading:

            if (
                found_first_heading
                and current_lines
            ):

                sections.append({
                    "chapter":
                        current_chapter,

                    "lines":
                        current_lines
                })

            current_chapter = heading

            current_lines = [
                line
            ]

            found_first_heading = True

        else:

            if found_first_heading:

                current_lines.append(
                    line
                )

    # Add final section
    if (
        found_first_heading
        and current_lines
    ):

        sections.append({
            "chapter":
                current_chapter,

            "lines":
                current_lines
        })

    # If no chapters were detected
    if not sections:

        sections.append({
            "chapter":
                default_chapter,

            "lines":
                lines
        })

    return sections


# ==========================================
# 10. NORMAL CHUNKING
# ==========================================

def chunk_chapter(
    lines,
    chunk_size=8,
    overlap=2
):

    chunks = []

    if not lines:
        return chunks

    # Prevent invalid step
    step = max(
        1,
        chunk_size - overlap
    )

    start = 0

    while start < len(lines):

        end = start + chunk_size

        chunk_lines = lines[
            start:end
        ]

        chunk = "\n".join(
            chunk_lines
        ).strip()

        if chunk:

            chunks.append(
                chunk
            )

        start += step

    return chunks


# ==========================================
# 11. QUESTION BANK CHUNKING
# ==========================================

QUESTION_START_PATTERN = re.compile(
    r"(?<!\d)(\d{1,3})\s*[\.\)]\s*"
)


def chunk_question_bank(text):
    """
    Convert a question-bank document into
    individual question strings.

    Handles:

        1. BFS
        2. DFS
        3) Monkey Banana
        10.Write Tower of Hanoi

    Also handles multiple questions on the
    same extracted PDF line.

    Garbage entries such as:

        10. 1

    are ignored.
    """

    # --------------------------------------
    # Clean input lines
    # --------------------------------------

    lines = []

    for raw_line in text.splitlines():

        cleaned = clean_line(
            raw_line
        )

        if cleaned:

            lines.append(
                cleaned
            )

    questions = []

    current_parts = []

    current_number = None

    # --------------------------------------
    # Helper to save current question
    # --------------------------------------

    def flush_current():

        nonlocal current_parts
        nonlocal current_number

        if not current_parts:
            return

        question = " ".join(
            current_parts
        ).strip()

        # Clean repeated spaces
        question = re.sub(
            r"\s+",
            " ",
            question
        ).strip()

        # ----------------------------------
        # Ignore empty / numeric-only garbage
        # ----------------------------------

        if not question:
            current_parts = []
            return

        if re.fullmatch(
            r"\d+",
            question
        ):

            current_parts = []
            return

        # Ignore a few obvious extraction
        # artifacts such as "1 2 3"
        if re.fullmatch(
            r"[\d\s]+",
            question
        ):

            current_parts = []
            return

        # ----------------------------------
        # Avoid duplicate questions
        # ----------------------------------

        existing = {
            q.lower()
            for q in questions
        }

        if question.lower() not in existing:

            questions.append(
                question
            )

        current_parts = []

    # --------------------------------------
    # Process lines
    # --------------------------------------

    for line in lines:

        matches = list(
            QUESTION_START_PATTERN.finditer(
                line
            )
        )

        # ----------------------------------
        # No numbered question on this line
        # ----------------------------------

        if not matches:

            if current_parts:

                current_parts.append(
                    line
                )

            continue

        # ----------------------------------
        # Process numbered questions
        # ----------------------------------

        for index, match in enumerate(matches):

            number = int(
                match.group(1)
            )

            # Content after this marker
            content_start = match.end()

            if index + 1 < len(matches):

                content_end = matches[
                    index + 1
                ].start()

            else:

                content_end = len(line)

            content = line[
                content_start:content_end
            ].strip()

            # ----------------------------------
            # Decide whether this is a real
            # question boundary.
            # ----------------------------------

            is_new_question = False

            if current_number is None:

                is_new_question = True

            elif number == current_number + 1:

                is_new_question = True

            elif (
                not current_parts
                and number == 1
            ):

                is_new_question = True

            # ----------------------------------
            # Start a new question
            # ----------------------------------

            if is_new_question:

                flush_current()

                current_number = number

                current_parts = []

                # ----------------------------------
                # Ignore numeric-only artifacts
                # ----------------------------------

                if content:

                    if not re.fullmatch(
                        r"\d+",
                        content
                    ):

                        current_parts.append(
                            content
                        )

            else:

                # ----------------------------------
                # Not a valid new question number.
                # Treat it as content.
                # ----------------------------------

                if content:

                    current_parts.append(
                        content
                    )

    # --------------------------------------
    # Save final question
    # --------------------------------------

    flush_current()

    return questions


# ==========================================
# 12. PROCESS PDF
# ==========================================

def process_pdf(
    file_path,
    subject="General",
    chapter="General"
):

    return process_file(
        file_path,
        subject,
        chapter
    )


# ==========================================
# 13. PROCESS ONE FILE
# ==========================================

def process_file(
    file_path,
    subject="General",
    chapter="General"
):

    # --------------------------------------
    # Extract text
    # --------------------------------------

    text = extract_text_from_file(
        file_path
    )

    filename = os.path.basename(
        file_path
    )

    # --------------------------------------
    # Determine document type
    # --------------------------------------

    document_type = get_document_type(
        file_path
    )

    documents = []

    # ======================================
    # QUESTION BANK
    # ======================================

    if document_type == "question_bank":

        questions = chunk_question_bank(
            text
        )

        for question in questions:

            documents.append({

                "text":
                    question,

                "metadata": {

                    "filename":
                        filename,

                    "subject":
                        subject,

                    "chapter":
                        chapter,

                    "content_type":
                        "Question Bank"
                }
            })

        return documents

    # ======================================
    # NORMAL DOCUMENT
    # ======================================

    sections = split_into_chapters(
        text=text,
        default_chapter=chapter,
        document_type=document_type
    )

    for section in sections:

        # ----------------------------------
        # Normalize Python chapter name
        # ----------------------------------

        if document_type == "chapter_style":

            chapter_name = (
                normalize_chapter_name(
                    section["chapter"]
                )
            )

        else:

            chapter_name = (
                section["chapter"]
            )

        # ----------------------------------
        # Create chunks
        # ----------------------------------

        chunks = chunk_chapter(
            section["lines"],
            chunk_size=8,
            overlap=2
        )

        # ----------------------------------
        # Create document objects
        # ----------------------------------

        for chunk in chunks:

            documents.append({

                "text":
                    chunk,

                "metadata": {

                    "filename":
                        filename,

                    "subject":
                        subject,

                    "chapter":
                        chapter_name
                }
            })

    return documents


# ==========================================
# 14. PROCESS MULTIPLE FILES
# ==========================================

def process_multiple_pdfs(
    documents
):

    all_documents = []

    for document in documents:

        processed = process_file(

            file_path=document[
                "file_path"
            ],

            subject=document.get(
                "subject",
                "General"
            ),

            chapter=document.get(
                "chapter",
                "General"
            )
        )

        all_documents.extend(
            processed
        )

    return all_documents